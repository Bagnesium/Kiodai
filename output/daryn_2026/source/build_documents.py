"""Build the editable manuscript and separate evidence note from their UTF-8 sources."""
import json,re
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path(__file__).resolve().parents[1]
PAGE_MAP=OUT/'data/page_map.json'
pagemap=json.loads(PAGE_MAP.read_text()) if PAGE_MAP.exists() else {}
TITLE='Kiodai: структурированная перспективная память языковых ИИ-агентов и методика её оценки'
HEADINGS=[]

def configure(d,note=False):
 s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(2.5);s.right_margin=s.top_margin=s.bottom_margin=Cm(2);s.footer_distance=Cm(.9);s.header_distance=Cm(.8);s.different_first_page_header_footer=True
 for style in d.styles:
  if style.type==1 or style.type==2:
   try:
    style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
   except AttributeError:pass
 normal=d.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.first_line_indent=Cm(.6);normal.paragraph_format.widow_control=True
 for name,size,before,after in [('Heading 1',14,11,7),('Heading 2',12,9,5),('Heading 3',12,7,4),('Title',20,0,14)]:
  st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.first_line_indent=0;st.paragraph_format.keep_with_next=True
 cap=d.styles['Caption'];cap.font.name='Times New Roman';cap.font.size=Pt(10.5);cap.font.italic=False;cap.font.bold=False;cap.paragraph_format.first_line_indent=0;cap.paragraph_format.line_spacing=1.03;cap.paragraph_format.space_after=Pt(7)
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=0
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);footer._p.append(f)
 for st in d.styles:
  for borders in list(st.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
  for fonts in st.element.iter(qn('w:rFonts')):
   for a in list(fonts.attrib):
    if 'Theme' in a:del fonts.attrib[a]
   for a in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+a),'Times New Roman')
 d.core_properties.title='Проверка источников Kiodai' if note else TITLE;d.core_properties.author='Беймжан Багдат / Bagdat Beimzhan';d.core_properties.subject='Перспективная память языковых ИИ-агентов';d.core_properties.language='ru-RU'
 lang=normal.element.get_or_add_rPr();el=OxmlElement('w:lang');el.set(qn('w:val'),'ru-RU');el.set(qn('w:eastAsia'),'kk-KZ');lang.append(el)

def bookmark(p,name):
 bid=str(int(name.rsplit('_',1)[1])+20)
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),bid);a.set(qn('w:name'),name)
 b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),bid);p._p.insert(0,a);p._p.append(b)

def bodypara(d,text,style=None,abstract=False,reference=False):
 p=d.add_paragraph(style=style);p.add_run(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if abstract:
  p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.05
 if reference or re.search(r'\S{38,}',text):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 if reference:
  p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.04
  for r in p.runs:r.font.size=Pt(11)
 return p

def titlepage(d):
 p=d.add_paragraph('Назарбаев Интеллектуальная школа\nестественно-математического направления\nрайона Нура города Астаны');p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(60)
 p=d.add_paragraph('НАУЧНЫЙ ПРОЕКТ');p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(17);p.runs[0].bold=True
 p=d.add_paragraph(TITLE,'Title');p.alignment=1;p.paragraph_format.space_after=Pt(32)
 p=d.add_paragraph('Направление II\n«Математическое моделирование экономических\nи социальных процессов»\nСекция: информатика');p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.line_spacing=1.18;p.paragraph_format.space_after=Pt(40)
 p=d.add_paragraph('Автор: Беймжан Багдат / Bagdat Beimzhan\nУченик 10 класса\n\nНаучный руководитель:\nОмаров Абылайхан Бауржанович');p.paragraph_format.first_line_indent=0;p.paragraph_format.left_indent=Cm(6);p.paragraph_format.space_after=Pt(38)
 p=d.add_paragraph('Астана, 2026');p.alignment=1;p.paragraph_format.first_line_indent=0
 d.add_page_break()

def toc(d,source):
 p=d.add_paragraph('ОГЛАВЛЕНИЕ');p.alignment=1;p.paragraph_format.first_line_indent=0;p.runs[0].bold=True;p.runs[0].font.size=Pt(14);p.paragraph_format.space_after=Pt(14)
 for line in source.splitlines():
  if not line.startswith('#'):continue
  level=len(line)-len(line.lstrip('#'));text=line[level:].strip()
  anchor='h_'+str(len(HEADINGS))
  HEADINGS.append((text,level,anchor))
  p=d.add_paragraph();p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
  p.paragraph_format.left_indent=Cm(.45 if level==2 else 0)
  p.paragraph_format.tab_stops.add_tab_stop(Cm(16.45),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');rpr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'23');rpr.append(sz);r.append(rpr);t=OxmlElement('w:t');t.text=text;r.append(t);link.append(r);p._p.append(link)
  p.add_run('\t'+str(pagemap.get(text,'0'))).font.size=Pt(11.5)
 d.add_page_break()

WIDTHS={
 1:[4.0,2.65,2.65,3.6,3.6],
 2:[4.9,11.6],
 3:[2.2,1.0,1.0,1.0,4.6,6.7],
 4:[2.7,2.1,4.1,3.7,3.9],
 5:[1.9,2.5,3.2,2.5,2.75,3.65],
 6:[1.85,2.0,2.0,3.0,2.2,3.6,1.85],
 7:[2.7,3.7,3.0,4.1,3.0],
}
# Widths below are normalized to the 16.5 cm text block, preserving relative allocations.
def table(d,rows,index):
 t=d.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=WIDTHS.get(index,[1]*len(rows[0]));widths=[w/sum(widths)*16.5 for w in widths]
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for ri,row in enumerate(rows):
  cells=t.rows[0].cells if ri==0 else t.add_row().cells
  for ci,(c,txt) in enumerate(zip(cells,row)):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
   p.paragraph_format.keep_with_next=(index<7 and ri<len(rows)-1)
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if index in [1,2,7] else WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(txt.replace('<br>','\n'));r.font.size=Pt(10 if index in [5,6,7] else 10.5);r.bold=ri==0
   pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for k,v in [('top','65'),('bottom','65'),('left','75'),('right','75')]:
    e=OxmlElement('w:'+k);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for k in ['top','left','bottom','right']:
    e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EFF2');pr.append(sh)
  trpr=t.rows[ri]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
 d.add_paragraph().paragraph_format.space_after=Pt(1)

def build(src,dest,note=False):
 d=Document();configure(d,note);source=src.read_text();lines=source.splitlines();i=0;abstract=False;refs=False;headidx=0;tabidx=0
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line=='@title':titlepage(d);continue
  if line=='@toc':toc(d,source);continue
  if line=='@abstracts':abstract=True;continue
  if line=='@body':abstract=False;d.add_page_break();continue
  if line=='@references':refs=True;d.add_page_break();continue
  if line.startswith('@figure '):
   fname,caption=line[8:].split(' | ',1)
   p=d.add_paragraph();p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
   r=p.add_run();r.add_picture(str(OUT/'figures'/fname),width=Cm(16.4))
   pic=p._p.xpath('.//wp:docPr')[0];pic.set('descr',caption)
   bodypara(d,caption,'Caption');continue
  if line.startswith('@equation '):
   p=d.add_paragraph();p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(3)
   math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=line[10:];r.append(t);math.append(r);p._p.append(math);continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));text=line[level:].strip()
   p=d.add_paragraph(text,'Heading '+str(min(level,3)));p.paragraph_format.first_line_indent=0
   if note and level==1:p.style='Title'
   if not note and text=='8. Заключение':p.paragraph_format.page_break_before=True
   if abstract:p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(3)
   if not note:
    anchor=HEADINGS[headidx][2];headidx+=1;bookmark(p,anchor)
   continue
  if line.startswith('|'):
   rows=[];block=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
   for x in block:
    if re.fullmatch(r'[|:\-\s]+',x):continue
    rows.append([c.strip() for c in x.strip('|').split('|')])
   tabidx+=1;table(d,rows,tabidx);continue
  if line.startswith('Таблица '):
   p=bodypara(d,line,'Caption');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4);continue
  bodypara(d,line,abstract=abstract,reference=refs)
 if note:
  for p in d.paragraphs:
   if p.style.name=='Normal':p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 d.save(dest)
 return d
if __name__=='__main__':
 main=build(OUT/'source/manuscript.md',OUT/'Kiodai_Daryn_2026.docx')
 build(OUT/'source/evidence_note.md',OUT/'Kiodai_Evidence_and_Verification.docx',True)
 (OUT/'data/headings.json').write_text(json.dumps(HEADINGS,ensure_ascii=False,indent=2))
 abstract_text=(OUT/'source/manuscript.md').read_text().split('@abstracts')[1].split('@body')[0]
 counts={}
 for block in abstract_text.split('### ')[1:]:
  title,body=block.split('\n',1);counts[title]=len(body.split())
 print('abstract word counts',counts)
 assert max(counts.values())<=250
 print('Main word count',len((OUT/'source/manuscript.md').read_text().split()))
